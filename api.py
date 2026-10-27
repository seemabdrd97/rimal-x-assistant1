import csv
import re
import threading
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path

from fastapi import FastAPI
from openpyxl import load_workbook
from pydantic import BaseModel

PROJECT_DIR = Path(__file__).resolve().parent
DOCUMENTS_DIR = PROJECT_DIR / "documents"
EXCEL_PATH = DOCUMENTS_DIR / "محتوى مساعد المعرفة.xlsx"
LOG_PATH = PROJECT_DIR / "logs" / "unanswered.csv"

PUBLIC_AUDIENCE = "الكل"
MIN_SCORE = 0.50

app = FastAPI(
    title="Rimal X Internal Assistant API",
    version="0.4.0",
)

_log_lock = threading.Lock()


class SearchRequest(BaseModel):
    question: str
    # الدور يأتي من العميل في النموذج الأولي فقط.
    # في النسخة الفعلية يجب أن يُستخرج من تسجيل الدخول على الخادم.
    role: str = PUBLIC_AUDIENCE


class SearchResult(BaseModel):
    status: str
    answer: str | None
    sources: list[dict]


STOP_WORDS = {
    "هل", "ما", "ماذا", "كيف", "كم", "متى", "من",
    "انا", "لي", "في", "على", "عن", "الى", "و", "او",
    "اقدر", "استطيع", "يمكنني", "ابي", "ابغى", "اريد",
    "هو", "هي", "هذا", "هذه", "منه", "عندي",
    "تم", "يتم", "كان", "كانت", "يكون",
    "مع", "بعد", "قبل", "اذا", "عند", "كل",
}

# توحيد بعض الصيغ المتقاربة، دون اعتبار كل الكلمات مترادفات
# ملاحظة: المفاتيح تُكتب بعد التطبيع (بدون همزات)
SYNONYMS = {
    "تنزل": "صرف",
    "ينزل": "صرف",
    "تنزلون": "صرف",
    "تصرف": "صرف",
    "تصرفون": "صرف",
    "يصرف": "صرف",
    "صرف": "صرف",
    "عمولتي": "عمولة",
    "عمولتك": "عمولة",
    "عمولته": "عمولة",
    "عمولتها": "عمولة",
    "العمولات": "عمولة",
    "عمولات": "عمولة",
    "العمولة": "عمولة",
    "العضويات": "عضوية",
    "عضويات": "عضوية",
    "العضوية": "عضوية",
    "بعت": "بيع",
    "ابيع": "بيع",
    "مواعيد": "موعد",
    "نسبه": "نسبة",
    "احسب": "حساب",
    "احسبها": "حساب",
    "حسابي": "حساب",
}

# لواصق شائعة (الترتيب مهم: الأطول أولًا)
PREFIXES = ("وال", "بال", "فال", "كال", "لل", "ال")


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", str(text)).casefold()
    text = re.sub(r"[\u064b-\u065f\u0670]", "", text)
    text = text.replace("ـ", "")
    text = re.sub(r"[إأآٱ]", "ا", text)
    text = text.replace("ى", "ي")
    text = re.sub(r"[^\w\s]", " ", text)
    return " ".join(text.split())


def keywords(text: str) -> set[str]:
    result = set()

    for word in normalize(text).split():
        if len(word) <= 1 or word in STOP_WORDS:
            continue

        word = SYNONYMS.get(word, word)

        for prefix in PREFIXES:
            if word.startswith(prefix) and len(word) >= len(prefix) + 3:
                word = word[len(prefix):]
                break

        # بعد إزالة اللاصقة قد تصبح الكلمة مرادفًا معروفًا
        word = SYNONYMS.get(word, word)
        result.add(word)

    return result


def get_cell(row, index):
    if index is None or index >= len(row):
        return ""

    value = row[index]

    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")

    return str(value).strip()


def split_alternatives(text: str) -> list[str]:
    return [
        part.strip()
        for part in re.split(r"[|؛;\n]+", text)
        if part.strip()
    ]


def similarity(question_words: set[str],
               candidate_words: set[str]) -> float:
    if not question_words or not candidate_words:
        return 0.0

    overlap = question_words & candidate_words

    if not overlap:
        return 0.0

    recall = len(overlap) / len(question_words)
    precision = len(overlap) / len(candidate_words)

    # نفضّل تطابق كلمات السؤال مع المرشح دون إهمال طول المرشح
    return 0.65 * recall + 0.35 * precision


def is_allowed(audience_cell: str, role: str) -> bool:
    """منع افتراضي: الخلية الفارغة تعني أن الصف غير متاح لأحد."""
    allowed = {normalize(a) for a in split_alternatives(audience_cell)}

    if not allowed:
        return False

    return (
        normalize(PUBLIC_AUDIENCE) in allowed
        or normalize(role) in allowed
    )


def redact(text: str) -> str:
    """إخفاء البيانات الشخصية الواضحة قبل حفظ السؤال في السجل."""
    text = re.sub(r"[\w.+-]+@[\w-]+\.[\w.-]+", "[بريد]", text)
    text = re.sub(r"\+?\d[\d\s\-]{6,}\d", "[رقم]", text)
    return " ".join(text.split())[:300]


def log_unanswered(question: str, role: str, reason: str) -> None:
    """يسجل السؤال دون هوية المستخدم، لتحسين المحتوى فقط."""
    try:
        LOG_PATH.parent.mkdir(parents=True, exist_ok=True)

        with _log_lock:
            is_new = not LOG_PATH.exists()

            with LOG_PATH.open("a", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)

                if is_new:
                    writer.writerow(
                        ["الوقت (UTC)", "السؤال", "الدور", "السبب"]
                    )

                writer.writerow([
                    datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M"),
                    redact(question),
                    redact(role),
                    reason,
                ])
    except Exception:
        # فشل التسجيل يجب ألا يعطل البحث
        pass


def no_answer(question: str, role: str, reason: str) -> SearchResult:
    log_unanswered(question, role, reason)
    return SearchResult(status="no_answer", answer=None, sources=[])


@app.get("/health")
def health():
    return {
        "status": "ok",
        "excel_exists": EXCEL_PATH.exists(),
    }


@app.post("/search", response_model=SearchResult)
def search_documents(request: SearchRequest):
    question = request.question
    role = request.role or PUBLIC_AUDIENCE
    question_words = keywords(question)

    if not question_words:
        return no_answer(question, role, "سؤال فارغ أو كلمات عامة فقط")

    if not EXCEL_PATH.exists():
        return no_answer(question, role, "ملف المحتوى غير موجود")

    matches = []
    restricted_best = 0.0
    workbook = None

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True,
        )

        sheet = workbook.active

        headers = [
            normalize(cell.value or "")
            for cell in sheet[1]
        ]

        column_map = {
            header: index
            for index, header in enumerate(headers)
            if header
        }

        def find_column(*names):
            for name in names:
                key = normalize(name)
                if key in column_map:
                    return column_map[key]
            return None

        original_col = find_column(
            "السؤال الأصلي", "السؤال"
        )
        alternatives_col = find_column(
            "الصيغ البديلة",
            "الصيغ البديلة (افصل بينها |)",
            "صيغ بديلة",
        )
        answer_col = find_column(
            "الإجابة",
            "الإجابة المعتمدة من المصدر",
            "الإجابة المعتمدة",
        )
        source_col = find_column(
            "المصدر", "رابط المصدر"
        )
        id_col = find_column(
            "معرّف السؤال", "معرف السؤال", "ID"
        )
        keywords_col = find_column(
            "الكلمات المفتاحية", "كلمات مفتاحية"
        )
        audience_col = find_column(
            "الفئة المسموحة", "الصلاحية", "الجمهور",
        )
        updated_col = find_column(
            "تاريخ آخر تحديث", "تاريخ التحديث", "آخر تحديث",
        )

        if original_col is None or answer_col is None:
            return no_answer(question, role, "أعمدة الملف غير مطابقة")

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            original = get_cell(row, original_col)
            answer = get_cell(row, answer_col)

            if not original or not answer:
                continue

            # الصلاحيات: إذا وُجد عمود الفئة المسموحة نطبقه بصرامة.
            # إذا لم يوجد العمود أصلًا نعمل بالوضع القديم (الكل مسموح).
            allowed = True
            if audience_col is not None:
                allowed = is_allowed(get_cell(row, audience_col), role)

            alternatives = get_cell(row, alternatives_col)
            extra_keywords = get_cell(row, keywords_col)

            candidates = [original]
            candidates.extend(split_alternatives(alternatives))

            best_score = 0.0

            for candidate in candidates:
                score = similarity(
                    question_words,
                    keywords(candidate),
                )
                best_score = max(best_score, score)

            # الكلمات المفتاحية تعزّز التطابق ولا تكفي وحدها
            if best_score > 0:
                keyword_overlap = question_words & keywords(extra_keywords)

                if keyword_overlap:
                    boost = 0.2 * len(keyword_overlap) / len(question_words)
                    best_score = min(1.0, best_score + boost)

            if best_score <= 0:
                continue

            if not allowed:
                restricted_best = max(restricted_best, best_score)
                continue

            matches.append({
                "score": best_score,
                "file": EXCEL_PATH.name,
                "line": row_number,
                "question": original,
                "answer": answer,
                "source": get_cell(row, source_col),
                "id": get_cell(row, id_col),
                "updated": get_cell(row, updated_col),
            })

    except Exception:
        return no_answer(question, role, "خطأ أثناء قراءة الملف")

    finally:
        if workbook is not None:
            workbook.close()

    if not matches:
        reason = (
            "يوجد تطابق لكنه خارج صلاحية الدور"
            if restricted_best >= MIN_SCORE
            else "لا يوجد تطابق"
        )
        return no_answer(question, role, reason)

    matches.sort(
        key=lambda item: item["score"],
        reverse=True,
    )

    best = matches[0]

    # لا نجيب إذا كان التطابق ضعيفًا
    if best["score"] < MIN_SCORE:
        reason = (
            "يوجد تطابق لكنه خارج صلاحية الدور"
            if restricted_best >= MIN_SCORE
            else "تطابق ضعيف"
        )
        return no_answer(question, role, reason)

    # نعرض فقط السؤال الذي اعتمدنا عليه كمصدر
    sources = [{
        "file": best["file"],
        "line": best["line"],
        "text": best["question"],
        "source": best["source"],
        "id": best["id"],
        "updated": best["updated"],
    }]

    return SearchResult(
        status="ok",
        answer=best["answer"],
        sources=sources,
    )
